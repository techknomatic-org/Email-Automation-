from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File
from sqlalchemy.orm import Session
from typing import List, Optional
from pydantic import BaseModel
from backend.app.core.database import get_db
from backend.app.services.rag_service import rag_service
from backend.app.models.knowledge import KnowledgeDocument, KnowledgeChunk

router = APIRouter(prefix="/knowledge", tags=["Company Knowledge RAG"])

class DocumentCreate(BaseModel):
    title: str
    category: Optional[str] = "General"
    content: str

class SearchQuery(BaseModel):
    query: str
    top_k: Optional[int] = 3

@router.get("")
def list_documents(db: Session = Depends(get_db)):
    try:
        return db.query(KnowledgeDocument).all()
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to list documents: {str(e)}")

@router.post("", status_code=status.HTTP_201_CREATED)
def create_document(payload: DocumentCreate, db: Session = Depends(get_db)):
    if not payload.title.strip():
        raise HTTPException(status_code=422, detail="Document title cannot be empty")
    if not payload.content.strip():
        raise HTTPException(status_code=422, detail="Document content cannot be empty")
    try:
        doc = rag_service.index_document(
            db,
            title=payload.title.strip(),
            category=payload.category or "General",
            content=payload.content.strip()
        )
        return {"id": doc.id, "title": doc.title, "category": doc.category, "message": "Document indexed successfully"}
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to index document: {str(e)}")

@router.post("/upload-file", status_code=status.HTTP_201_CREATED)
async def upload_file(
    file: UploadFile = File(...),
    category: Optional[str] = "General",
    db: Session = Depends(get_db)
):
    filename = file.filename
    ext = filename.split('.')[-1].lower() if '.' in filename else ''
    
    contents = await file.read()
    text_content = ""
    
    try:
        if ext in ['txt', 'md', 'markdown']:
            text_content = contents.decode('utf-8', errors='ignore')
        elif ext == 'csv':
            text_content = contents.decode('utf-8', errors='ignore')
            import csv
            import io
            f = io.StringIO(text_content)
            reader = csv.reader(f)
            rows = list(reader)
            if rows:
                text_content = "\n".join([" | ".join(row) for row in rows])
        elif ext == 'pdf':
            import io
            from pypdf import PdfReader
            pdf = PdfReader(io.BytesIO(contents))
            pages = []
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    pages.append(text)
            text_content = "\n\n".join(pages)
            if not text_content.strip():
                raise HTTPException(status_code=400, detail="PDF contains no extractable text (it might be scanned).")
        elif ext in ['docx', 'doc']:
            import io
            from docx import Document
            doc = Document(io.BytesIO(contents))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            text_content = "\n\n".join(paragraphs)
        elif ext in ['xlsx', 'xls']:
            import io
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
            sheet_texts = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows_text = []
                for row in ws.iter_rows(values_only=True):
                    if any(v is not None for v in row):
                        row_vals = [str(v) if v is not None else "" for v in row]
                        rows_text.append(" | ".join(row_vals))
                if rows_text:
                    sheet_texts.append(f"### Sheet: {sheet_name}\n" + "\n".join(rows_text))
            text_content = "\n\n".join(sheet_texts)
        else:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported file format: .{ext}. Please upload .txt, .md, .csv, .pdf, .docx, or .xlsx files."
            )
            
        if not text_content.strip():
            raise HTTPException(status_code=400, detail="The uploaded file contains no extractable text.")
            
        text_content = text_content.replace('\x00', '')
        title = filename.rsplit('.', 1)[0] if '.' in filename else filename
        
        doc = rag_service.index_document(
            db,
            title=title,
            category=category or "General",
            content=text_content
        )
        
        return {
            "id": doc.id,
            "title": doc.title,
            "category": doc.category,
            "message": "File parsed and indexed successfully"
        }
    except HTTPException as e:
        raise e
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail=f"Failed to process and index file: {str(e)}"
        )

@router.post("/search")
def search_knowledge(payload: SearchQuery, db: Session = Depends(get_db)):
    if not payload.query.strip():
        raise HTTPException(status_code=422, detail="Search query cannot be empty")
    try:
        results = rag_service.search_knowledge(db, query=payload.query.strip(), top_k=payload.top_k or 3)
        return {"query": payload.query, "results": results, "count": len(results)}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Search failed: {str(e)}")

@router.delete("/{doc_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_document(doc_id: int, db: Session = Depends(get_db)):
    doc = db.query(KnowledgeDocument).filter(KnowledgeDocument.id == doc_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    try:
        db.delete(doc)
        db.commit()
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to delete: {str(e)}")
    return None
